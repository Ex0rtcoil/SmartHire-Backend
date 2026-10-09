from pathlib import Path
from io import BytesIO

from flask import Flask, jsonify, request, send_from_directory
from pypdf import PdfReader
from docx import Document
from werkzeug.utils import secure_filename

from database import init_db, save_result
from classifier import classify_resume

BASE_DIR = Path(__file__).resolve().parent.parent
MAX_SIZE = 5 * 1024 * 1024
ALLOWED_EXTENSIONS = {".pdf", ".docx"}

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 30 * 1024 * 1024

init_db()


@app.get("/")
def home():
    return send_from_directory(BASE_DIR, "index.html")


@app.get("/webstyle.css")
def stylesheet():
    return send_from_directory(BASE_DIR, "webstyle.css")


@app.get("/script.js")
def javascript():
    return send_from_directory(BASE_DIR, "script.js")


def extract_text(file_bytes, extension):
    if extension == ".pdf":
        reader = PdfReader(BytesIO(file_bytes))
        return "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )

    document = Document(BytesIO(file_bytes))
    paragraphs = [p.text for p in document.paragraphs]

    for table in document.tables:
        for row in table.rows:
            paragraphs.append(
                " ".join(cell.text for cell in row.cells)
            )

    return "\n".join(paragraphs)


@app.post("/api/classify")
def classify():
    uploaded_files = request.files.getlist("files")

    if not uploaded_files or all(
        not f.filename for f in uploaded_files
    ):
        return jsonify({"error": "Please upload at least one resume."}), 400

    results = []

    for uploaded in uploaded_files:
        if not uploaded.filename:
            continue

        filename = secure_filename(uploaded.filename)
        extension = Path(filename).suffix.lower()

        if extension not in ALLOWED_EXTENSIONS:
            results.append({
                "file": filename,
                "error": "Only PDF and DOCX files are supported."
            })
            continue

        file_bytes = uploaded.read()

        if len(file_bytes) > MAX_SIZE:
            results.append({
                "file": filename,
                "error": "This file exceeds the 5 MB limit."
            })
            continue

        try:
            text = extract_text(file_bytes, extension)

            if not text.strip():
                results.append({
                    "file": filename,
                    "error": (
                        "No readable text was found. "
                        "Scanned PDFs may require OCR."
                    )
                })
                continue

            result = classify_resume(text)
            result["file"] = filename

            save_result(filename, result)
            results.append(result)

        except Exception:
            app.logger.exception(
                "Could not process uploaded resume: %s",
                filename
            )
            results.append({
                "file": filename,
                "error": (
                    "The file could not be processed. "
                    "Check that it is a valid PDF or DOCX file."
                )
            })

    return jsonify(results)


@app.errorhandler(413)
def too_large(_error):
    return jsonify({
        "error": "The total upload is too large. Try fewer files."
    }), 413


if __name__ == "__main__":
    app.run(debug=True)